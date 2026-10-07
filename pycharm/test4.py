import os.path
import random
from random import randint

import numpy as np
import openpyxl

def maopao(arr):
    n=len(arr)
    for i in range(n):
        for j in range(n-i-1):
            if arr[j]>arr[j+1]:
               arr[j],arr[j+1]=arr[j+1],arr[j]
    return arr
num=random.sample(range(0,100),5)
print(f"需要排序的数组为：",num)
newnum=maopao(num)
print('排序(升序)后的数组为：',*newnum,sep=" ")

# 处理表格
import openpyxl
import pandas as pd
import numpy as np
path=r'D:\pycharm\jg.xlsx'
df=pd.read_excel(path,header=1)
#三目运算，“成立时结果或者要进行的操作” if 判断事件 else 不成立的结果或操作“
df=df.map(lambda x: x.replace(" ","") if isinstance(x,str) else x).replace("",np.nan)
#有loc 修改表格对应的内容
df.loc[df['申办流水号'] == '520525163424001223010500001',['申请人','受理人']] = ['何秀珍贵','龙大纲']
print(f"{df['收件时间'].dtype}")
df['收件时间']=pd.to_datetime(df['收件时间'],format='%Y-%m-%d%H:%M:%S.%f',errors='coerce')
print(f"{df['收件时间'].dtype}")
df_re=df[df['收件时间'].between('2021-01-01','2025-12-31')] #按时间筛选
grouped=df.groupby('审批单位')
re=list(grouped.groups.keys())
try:
   with pd.ExcelWriter(path,engine='openpyxl',mode='a') as writer:
        if "原始数据" in writer.book.sheetnames:
           writer.book.remove(writer.book["原始数据"])
        df.to_excel(writer,sheet_name="原始数据",index=True)
        for dz,gr in grouped:
            if dz in writer.book.sheetnames:
               writer.book.remove(writer.book[dz])
            gr.to_excel(writer,sheet_name=dz,index=False)
   print(f"✅ 分组完成，分组结果：{','.join(re)},共{len(re)}组。")
except PermissionError:
    print(f'❌ 错误：文件 {path} 正在被使用，请关闭 Excel 文件后重试。')

path=r"C:\Users\qlok\Desktop\心怀感恩.docx"
from docx import Document
txt=Document(path)
txt1=[p.text for p in txt.paragraphs]
ntx=Document(path)
tar=ntx.paragraphs[3]
tj=('  纳雍县交通运输局,纳雍县人力资源和社会保障局,纳雍县公安局,纳雍县利园社区,\
           纳雍县医疗保障局,纳雍县卫生健康局,纳雍县市场监督管理局,纳雍县林业局,纳雍县民族宗教事务局')
tar.insert_paragraph_before(tj)
nm='test.docx'
ntx.save(nm)
pt=os.path.join(os.getcwd(),nm) #os.getcwd()获取路径
print(f"文件保存在：{pt}")

